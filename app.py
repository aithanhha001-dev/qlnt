import calendar
import hashlib
import io
import os
import re
import warnings
from contextlib import contextmanager
from datetime import date, datetime

import pandas as pd
import streamlit as st

# ==========================================
# 1. CẤU HÌNH TRANG & HÀM TIỆN ÍCH
# ==========================================
st.set_page_config(page_title="Quản lý nhân sự", layout="wide", page_icon="🏫")

# Ẩn thanh công cụ của Streamlit
st.markdown("""
<style>
[data-testid="stToolbar"], [data-testid="stMainMenu"], [data-testid="stAppDeployButton"],
[data-testid="stDeployButton"], .stDeployButton, #MainMenu, footer,
[data-testid="stDecoration"], [data-testid="stStatusWidget"] { display: none !important; visibility: hidden !important; }
</style>
""", unsafe_allow_html=True)

_ST_VER = tuple(int(x) for x in re.findall(r"\d+", st.__version__)[:2])
W = {"width": "stretch"} if _ST_VER >= (1, 50) else {"use_container_width": True}

for _k, _v in {"logged_in": False, "username": "", "role": "", "ma_vien_chuc": "", "flash": ""}.items():
    st.session_state.setdefault(_k, _v)

MIN_DATE = date(1940, 1, 1)
MAX_DATE = date(2040, 12, 31)
DATE_FORMAT = "DD/MM/YYYY"
YEARS = list(range(2020, 2041))
GIOI_TINH = ["Nam", "Nữ"]
TRINH_DO = ["Đại học", "Cao đẳng", "Trung cấp", "Thạc sĩ", "Tiến sĩ", "Khác"]
HANG_CDNN = ["Hạng I", "Hạng II", "Hạng III", "Hạng IV", "Khác"]
TRANG_THAI_CC = ["X", "P", "K", "VR", "CT", "O"]
LUONG_CO_SO = 2530000
CHUC_VU_LIST = ["Hiệu trưởng", "Phó hiệu trưởng", "Tổ trưởng CM", "Tổ phó CM", "Giáo viên",
                "Nhân viên y tế", "Nhân viên Hành chính", "Nhân viên khác"]
CHUC_VU_CU = {"Phó Hiệu trưởng": "Phó hiệu trưởng", "Tổ trưởng": "Tổ trưởng CM", "Nhân viên": "Nhân viên khác"}

def make_hashes(password): return hashlib.sha256(str.encode(password)).hexdigest()
def check_hashes(password, hashed_text): return make_hashes(password) == hashed_text

# ==========================================
# 2. MODULE KẾT NỐI CƠ SỞ DỮ LIỆU ĐÁM MÂY
# ==========================================
try:
    HAS_SECRETS = "DATABASE_URL" in st.secrets
except Exception:
    HAS_SECRETS = False
USE_POSTGRES = HAS_SECRETS or "DATABASE_URL" in os.environ

if USE_POSTGRES:
    import psycopg2
    from psycopg2 import IntegrityError
    DB_URL = st.secrets.get("DATABASE_URL", os.environ.get("DATABASE_URL", ""))
    
    @contextmanager
    def db():
        conn = psycopg2.connect(DB_URL)
        try:
            yield conn
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            conn.close()

    def q(sql, params=()):
        sql = sql.replace('?', '%s')
        sql = sql.replace("CAST(strftime('%d', ngay_cham_cong) AS INTEGER)", "CAST(EXTRACT(DAY FROM ngay_cham_cong) AS INTEGER)")
        with db() as conn:
            with warnings.catch_warnings():
                warnings.simplefilter('ignore', UserWarning)
                return pd.read_sql_query(sql, conn, params=params)

    def run(sql, params=()):
        sql = sql.replace('?', '%s')
        with db() as conn:
            with conn.cursor() as cur:
                cur.execute(sql, params)

    def run_many(sql, records):
        sql = sql.replace('?', '%s')
        with db() as conn:
            with conn.cursor() as cur:
                cur.executemany(sql, records)
else:
    import sqlite3
    from sqlite3 import IntegrityError
    DB_NAME = "quanly_nhatruong.db"

    @contextmanager
    def db():
        conn = sqlite3.connect(DB_NAME)
        try:
            yield conn
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        finally:
            conn.close()

    def q(sql, params=()):
        with db() as conn: return pd.read_sql_query(sql, conn, params=params)

    def run(sql, params=()):
        with db() as conn: conn.execute(sql, params)

    def run_many(sql, records):
        with db() as conn: conn.executemany(sql, records)

def flash_rerun(msg, *clear_keys):
    st.session_state["flash"] = msg
    for k in clear_keys: st.session_state.pop(k, None)
    st.rerun()

# ---------- Hàm Ngày tháng / định dạng ----------
def date_or_none(val):
    if val is None: return None
    try:
        if pd.isna(val): return None
    except (TypeError, ValueError): pass
    s = str(val).strip()
    if s in ("", "NaT", "None", "nan", "<NA>"): return None
    try:
        if isinstance(val, datetime): d = val.date()
        elif isinstance(val, date): d = val
        elif re.match(r"^\d{4}-\d{2}-\d{2}", s): d = pd.to_datetime(s).date()
        else: d = pd.to_datetime(s, dayfirst=True).date()
    except Exception: return None
    return min(max(d, MIN_DATE), MAX_DATE)

def date_or_today(val): return date_or_none(val) or date.today()
def sql_date(d): return d.strftime("%Y-%m-%d") if d else None
def format_date_vn(val):
    d = date_or_none(val)
    return d.strftime("%d/%m/%Y") if d else ""

def safe_str(val):
    try: return "" if val is None or pd.isna(val) else str(val)
    except (TypeError, ValueError): return str(val)

def pick(options, value, default=0): return options.index(value) if value in options else default

# ---------- Hàm Excel ----------
def _autofit(ws, df):
    from openpyxl.utils import get_column_letter
    for i, col in enumerate(df.columns, start=1):
        width = max([len(str(col))] + [len(str(v)) for v in df[col].head(500)]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(max(width, 8), 40)
    ws.freeze_panes = "A2"

def to_excel_sheets(sheets):
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, index=False, sheet_name=name[:31])
            _autofit(writer.sheets[name[:31]], df)
    return output.getvalue()

def to_excel_bytes(df, sheet_name="DanhSach"): return to_excel_sheets({sheet_name: df})

def quarter_of(ts): return (ts.month - 1) // 3 + 1

def filter_quarter(df, col, year, quarter):
    mask = df[col].apply(lambda t: pd.notna(t) and t.year == year and quarter_of(t) == quarter)
    return df[mask].sort_values(col)

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

def quarter_export_ui(df, date_col, cols, title, file_prefix, key):
    st.subheader(title)
    today = date.today()
    c1, c2 = st.columns(2)
    nam_x = c1.number_input("Chọn Năm", min_value=2000, max_value=2100, value=today.year, step=1, key=f"{key}_nam")
    quy_x = c2.selectbox("Chọn Quý", [1, 2, 3, 4], index=quarter_of(today) - 1,
                         format_func=lambda x: f"Quý {x} (tháng {3 * x - 2}-{3 * x})", key=f"{key}_quy")
    if df.empty:
        st.info("Chưa có dữ liệu.")
        return

    def prep(sub):
        out = sub[cols].copy()
        out.insert(0, "STT", range(1, len(out) + 1))
        return out

    df_q = prep(filter_quarter(df, date_col, int(nam_x), int(quy_x)))
    if df_q.empty:
        st.info(f"Không có cán bộ đến hạn trong Quý {quy_x}/{int(nam_x)}.")
    else:
        st.dataframe(df_q, hide_index=True, **W)
        st.download_button(f"📥 Xuất Excel Quý {quy_x}/{int(nam_x)}", data=to_excel_bytes(df_q, f"Quy{quy_x}_{int(nam_x)}"),
                           mime=XLSX_MIME, file_name=f"{file_prefix}_Q{quy_x}_{int(nam_x)}.xlsx", key=f"{key}_dl_q")
    all_sheets = {f"Quy{i}": prep(filter_quarter(df, date_col, int(nam_x), i)) for i in range(1, 5)}
    if any(not v.empty for v in all_sheets.values()):
        st.download_button(f"📥 Xuất cả năm {int(nam_x)} (4 sheet theo quý)", data=to_excel_sheets(all_sheets),
                           mime=XLSX_MIME, file_name=f"{file_prefix}_Nam{int(nam_x)}.xlsx", key=f"{key}_dl_year")

def generate_excel_template():
    cols = ["Mã VC", "Họ và tên", "Giới tính", "Ngày sinh", "Dân tộc", "Ngày tuyển dụng",
            "Ngày vào Đảng", "Trình độ", "Chuyên ngành", "Ghi chú"]
    return to_excel_bytes(pd.DataFrame(columns=cols))

def style_cham_cong(val):
    return {
        "P": "color: red; font-weight: bold;",
        "VR": "color: blue; font-weight: bold;",
        "K": "color: white; background-color: darkred; font-weight: bold;",
        "CT": "color: purple; font-weight: bold;",
        "O": "color: #ff9800; font-weight: bold;",
        "X": "color: #4caf50;",
    }.get(val, "")

def style_days(df, day_cols):
    s = df.style
    fn = s.map if hasattr(s, "map") else s.applymap
    return fn(style_cham_cong, subset=day_cols)

# ==========================================
# 3. KHỞI TẠO BẢNG TỰ ĐỘNG
# ==========================================
VC_EXTRA_COLS = {
    "he_so_luong": "REAL", "phu_cap": "REAL", "ngay_huong_luong": "DATE", "chu_ky_nang_luong": "INTEGER",
    "muc_tham_nien": "REAL", "ngay_huong_tham_nien": "DATE", "chu_ky_nang_tham_nien": "INTEGER",
}

@st.cache_resource(show_spinner=False)
def init_db():
    try:
        with db() as conn:
            if USE_POSTGRES:
                with conn.cursor() as c:
                    c.execute("""CREATE TABLE IF NOT EXISTS users (
                                 id SERIAL PRIMARY KEY, username TEXT UNIQUE NOT NULL, password TEXT NOT NULL, 
                                 role TEXT NOT NULL, ma_vien_chuc TEXT)""")
                    c.execute("""CREATE TABLE IF NOT EXISTS chuc_vu (id SERIAL PRIMARY KEY, ten_chuc_vu TEXT NOT NULL)""")
                    c.execute("""CREATE TABLE IF NOT EXISTS vien_chuc (
                                    id SERIAL PRIMARY KEY, ma_vien_chuc TEXT UNIQUE NOT NULL, ho_ten TEXT NOT NULL,
                                    gioi_tinh TEXT, ngay_sinh DATE, dan_toc TEXT, chuc_vu_id INTEGER, ngay_tuyen_dung DATE,
                                    ngay_vao_dang DATE, hang_chuc_danh TEXT, trinh_do TEXT, ten_truong_dao_tao TEXT,
                                    chuyen_nganh TEXT, hinh_thuc_dao_tao TEXT, thac_si_chuyen_nganh TEXT, quan_ly_nha_nuoc TEXT,
                                    ly_luan_chinh_tri TEXT, tin_hoc TEXT, ngoai_ngu TEXT, chung_chi_khac TEXT, ghi_chu TEXT,
                                    he_so_luong REAL, phu_cap REAL, ngay_huong_luong DATE, chu_ky_nang_luong INTEGER,
                                    muc_tham_nien REAL, ngay_huong_tham_nien DATE, chu_ky_nang_tham_nien INTEGER,
                                    FOREIGN KEY(chuc_vu_id) REFERENCES chuc_vu(id))""")
                    c.execute("""CREATE TABLE IF NOT EXISTS cham_cong (
                                    id SERIAL PRIMARY KEY, vien_chuc_id INTEGER NOT NULL,
                                    ngay_cham_cong DATE NOT NULL, trang_thai TEXT DEFAULT 'X', ghi_chu TEXT,
                                    UNIQUE(vien_chuc_id, ngay_cham_cong))""")
                    c.execute("""CREATE TABLE IF NOT EXISTS khen_thuong (
                                    id SERIAL PRIMARY KEY, vien_chuc_id INTEGER NOT NULL,
                                    ngay_thang DATE, thanh_tich TEXT, danh_hieu TEXT, hinh_thuc TEXT,
                                    cap_khen TEXT, nam_khen INTEGER, so_quyet_dinh TEXT,
                                    FOREIGN KEY(vien_chuc_id) REFERENCES vien_chuc(id))""")
                    c.execute("""CREATE TABLE IF NOT EXISTS vu_viec (
                                    id SERIAL PRIMARY KEY, ngay_thang DATE, noi_dung TEXT, phuong_an TEXT)""")

                    def add_missing(table, cols):
                        c.execute("SELECT column_name FROM information_schema.columns WHERE table_name = %s", (table,))
                        existing = [r[0] for r in c.fetchall()]
                        for name, typ in cols.items():
                            if name not in existing:
                                c.execute(f"ALTER TABLE {table} ADD COLUMN {name} {typ}")
                    
                    add_missing("users", {"ma_vien_chuc": "TEXT"})
                    add_missing("vien_chuc", VC_EXTRA_COLS)

                    c.execute("""INSERT INTO users (username, password, role, ma_vien_chuc) 
                                 VALUES (%s, %s, %s, '') 
                                 ON CONFLICT (username) DO NOTHING""",
                              ("admin", make_hashes("admin123"), "admin"))
                    
                    for old, new in CHUC_VU_CU.items():
                        c.execute("SELECT 1 FROM chuc_vu WHERE ten_chuc_vu=%s", (new,))
                        if not c.fetchone():
                            c.execute("UPDATE chuc_vu SET ten_chuc_vu=%s WHERE ten_chuc_vu=%s", (new, old))
                    for ten in CHUC_VU_LIST:
                        c.execute("SELECT 1 FROM chuc_vu WHERE ten_chuc_vu=%s", (ten,))
                        if not c.fetchone():
                            c.execute("INSERT INTO chuc_vu (ten_chuc_vu) VALUES (%s)", (ten,))
                    
                    placeholders = ','.join(['%s'] * len(CHUC_VU_LIST))
                    c.execute(f"""DELETE FROM chuc_vu WHERE ten_chuc_vu NOT IN ({placeholders})
                                  AND id NOT IN (SELECT chuc_vu_id FROM vien_chuc WHERE chuc_vu_id IS NOT NULL)""", CHUC_VU_LIST)
            else:
                c = conn.cursor()
                c.execute('''CREATE TABLE IF NOT EXISTS users (id INTEGER PRIMARY KEY AUTOINCREMENT,
                             username TEXT UNIQUE NOT NULL, password TEXT NOT NULL, role TEXT NOT NULL, ma_vien_chuc TEXT)''')
                c.execute('''CREATE TABLE IF NOT EXISTS chuc_vu (id INTEGER PRIMARY KEY AUTOINCREMENT, ten_chuc_vu TEXT NOT NULL)''')
                c.execute('''CREATE TABLE IF NOT EXISTS vien_chuc (
                                id INTEGER PRIMARY KEY AUTOINCREMENT, ma_vien_chuc TEXT UNIQUE NOT NULL, ho_ten TEXT NOT NULL,
                                gioi_tinh TEXT, ngay_sinh DATE, dan_toc TEXT, chuc_vu_id INTEGER, ngay_tuyen_dung DATE,
                                ngay_vao_dang DATE, hang_chuc_danh TEXT, trinh_do TEXT, ten_truong_dao_tao TEXT,
                                chuyen_nganh TEXT, hinh_thuc_dao_tao TEXT, thac_si_chuyen_nganh TEXT, quan_ly_nha_nuoc TEXT,
                                ly_luan_chinh_tri TEXT, tin_hoc TEXT, ngoai_ngu TEXT, chung_chi_khac TEXT, ghi_chu TEXT,
                                he_so_luong REAL, phu_cap REAL, ngay_huong_luong DATE, chu_ky_nang_luong INTEGER,
                                muc_tham_nien REAL, ngay_huong_tham_nien DATE, chu_ky_nang_tham_nien INTEGER,
                                FOREIGN KEY(chuc_vu_id) REFERENCES chuc_vu(id))''')
                c.execute('''CREATE TABLE IF NOT EXISTS cham_cong (
                                id INTEGER PRIMARY KEY AUTOINCREMENT, vien_chuc_id INTEGER NOT NULL,
                                ngay_cham_cong DATE NOT NULL, trang_thai TEXT DEFAULT 'X', ghi_chu TEXT,
                                UNIQUE(vien_chuc_id, ngay_cham_cong))''')
                c.execute('''CREATE TABLE IF NOT EXISTS khen_thuong (
                                id INTEGER PRIMARY KEY AUTOINCREMENT, vien_chuc_id INTEGER NOT NULL,
                                ngay_thang DATE, thanh_tich TEXT, danh_hieu TEXT, hinh_thuc TEXT,
                                cap_khen TEXT, nam_khen INTEGER, so_quyet_dinh TEXT,
                                FOREIGN KEY(vien_chuc_id) REFERENCES vien_chuc(id))''')
                c.execute('''CREATE TABLE IF NOT EXISTS vu_viec (
                                id INTEGER PRIMARY KEY AUTOINCREMENT, ngay_thang DATE, noi_dung TEXT, phuong_an TEXT)''')

                def add_missing(table, cols):
                    existing = [r[1] for r in c.execute(f"PRAGMA table_info({table})").fetchall()]
                    for name, typ in cols.items():
                        if name not in existing: c.execute(f"ALTER TABLE {table} ADD COLUMN {name} {typ}")

                add_missing("users", {"ma_vien_chuc": "TEXT"})
                add_missing("vien_chuc", VC_EXTRA_COLS)

                c.execute("INSERT OR IGNORE INTO users (username, password, role, ma_vien_chuc) VALUES (?, ?, ?, '')",
                          ("admin", make_hashes("admin123"), "admin"))
                for old, new in CHUC_VU_CU.items():
                    if not c.execute("SELECT 1 FROM chuc_vu WHERE ten_chuc_vu=?", (new,)).fetchone():
                        c.execute("UPDATE chuc_vu SET ten_chuc_vu=? WHERE ten_chuc_vu=?", (new, old))
                for ten in CHUC_VU_LIST:
                    if not c.execute("SELECT 1 FROM chuc_vu WHERE ten_chuc_vu=?", (ten,)).fetchone():
                        c.execute("INSERT INTO chuc_vu (ten_chuc_vu) VALUES (?)", (ten,))
                c.execute(f"""DELETE FROM chuc_vu WHERE ten_chuc_vu NOT IN ({','.join('?' * len(CHUC_VU_LIST))})
                              AND id NOT IN (SELECT chuc_vu_id FROM vien_chuc WHERE chuc_vu_id IS NOT NULL)""", CHUC_VU_LIST)
    except Exception as e:
        print(f"Lỗi khởi tạo DB: {e}")

init_db()

# ==========================================
# 4. MÀN HÌNH ĐĂNG NHẬP
# ==========================================
def logout():
    st.session_state.update({"logged_in": False, "username": "", "role": "", "ma_vien_chuc": "", "flash": ""})
    st.rerun()

if not st.session_state["logged_in"]:
    st.markdown("<h1 style='text-align: center; color: #1E88E5; margin-top: 50px;'>🏫 HỆ THỐNG QUẢN LÝ NHÂN SỰ</h1>", unsafe_allow_html=True)
    st.markdown("---")
    _, col2, _ = st.columns([1, 1, 1])
    with col2:
        with st.form("login_form"):
            st.subheader("🔒 Đăng nhập")
            username = st.text_input("Tài khoản").strip()
            password = st.text_input("Mật khẩu", type="password")
            submitted = st.form_submit_button("🚀 Đăng nhập", **W)
        if submitted:
            res = q("SELECT password, role, ma_vien_chuc FROM users WHERE username=?", (username,))
            if not res.empty and check_hashes(password, res.iloc[0]["password"]):
                st.session_state.update({"logged_in": True, "username": username, "role": res.iloc[0]["role"], "ma_vien_chuc": safe_str(res.iloc[0]["ma_vien_chuc"])})
                st.rerun()
            else:
                st.error("❌ Tài khoản hoặc mật khẩu không chính xác!")
    st.stop()

is_admin = st.session_state["role"] == "admin"
current_ma_vc = st.session_state["ma_vien_chuc"]

if not is_admin and not current_ma_vc:
    st.error("Tài khoản của bạn chưa được liên kết với hồ sơ Cán bộ nào. Vui lòng liên hệ Quản trị viên.")
    if st.button("Đăng xuất"): logout()
    st.stop()

def get_current_vc_id():
    if is_admin: return None
    res = q("SELECT id FROM vien_chuc WHERE ma_vien_chuc = ?", (current_ma_vc,))
    return int(res.iloc[0]["id"]) if not res.empty else None

def scope_where(alias="vc", prefix="WHERE"):
    if is_admin: return "", ()
    return f"{prefix} {alias}.ma_vien_chuc = ?", (current_ma_vc,)

def delete_vien_chuc(ids):
    ids = [int(i) for i in ids]
    if not ids: return
    ph = ",".join("?" * len(ids))
    with db() as conn:
        if USE_POSTGRES:
            cur = conn.cursor()
            cur.execute(f"SELECT ma_vien_chuc FROM vien_chuc WHERE id IN ({ph})".replace('?', '%s'), ids)
            codes = [r[0] for r in cur.fetchall()]
            cur.execute(f"DELETE FROM cham_cong WHERE vien_chuc_id IN ({ph})".replace('?', '%s'), ids)
            cur.execute(f"DELETE FROM khen_thuong WHERE vien_chuc_id IN ({ph})".replace('?', '%s'), ids)
            cur.execute(f"DELETE FROM vien_chuc WHERE id IN ({ph})".replace('?', '%s'), ids)
            if codes:
                cur.execute(f"UPDATE users SET ma_vien_chuc='' WHERE ma_vien_chuc IN ({','.join('?' * len(codes))})".replace('?', '%s'), codes)
            cur.close()
        else:
            codes = [r[0] for r in conn.execute(f"SELECT ma_vien_chuc FROM vien_chuc WHERE id IN ({ph})", ids)]
            conn.execute(f"DELETE FROM cham_cong WHERE vien_chuc_id IN ({ph})", ids)
            conn.execute(f"DELETE FROM khen_thuong WHERE vien_chuc_id IN ({ph})", ids)
            conn.execute(f"DELETE FROM vien_chuc WHERE id IN ({ph})", ids)
            if codes:
                conn.execute(f"UPDATE users SET ma_vien_chuc='' WHERE ma_vien_chuc IN ({','.join('?' * len(codes))})", codes)

df_chuc_vu = q("SELECT * FROM chuc_vu")
_order = lambda t: CHUC_VU_LIST.index(t) if t in CHUC_VU_LIST else len(CHUC_VU_LIST)
df_chuc_vu = df_chuc_vu.assign(_o=df_chuc_vu["ten_chuc_vu"].map(_order)).sort_values(["_o", "id"])
CV_NAMES = dict(zip(df_chuc_vu["id"].astype(int).tolist(), df_chuc_vu["ten_chuc_vu"].tolist()))

st.sidebar.title("🏫 Menu Chức Năng")
st.sidebar.info(f"👤 Cán bộ: **{st.session_state['username'].upper()}**\n\n🔑 Quyền: **{'Admin' if is_admin else 'Giáo viên'}**\n\n🌐 Mạng: **{'Postgres Đám mây' if USE_POSTGRES else 'SQLite Local'}**")

menu_options = ["📊 Tổng quan", "👥 Quản lý Nhân sự", "⏱️ Chấm công", "💰 Quản lý Lương", "⏳ Quản lý Thâm niên", "🏆 Thi đua - Khen thưởng", "🚨 Quản lý Vụ việc"]
if is_admin: menu_options.append("⚙️ Quản lý Tài khoản")
menu = st.sidebar.radio("Điều hướng", menu_options)
st.sidebar.markdown("---")
if st.sidebar.button("🚪 Đăng xuất", **W): logout()

if st.session_state.get("flash"): st.success(st.session_state.pop("flash"))

# ==========================================
# CÁC MODULE CHỨC NĂNG
# ==========================================
def vc_fields(info=None):
    edit = info is not None
    g = (lambda k: safe_str(info[k])) if edit else (lambda k: "")
    d = (lambda k: date_or_none(info[k])) if edit else (lambda k: None)

    st.markdown("**1. Thông tin cơ bản**")
    c1, c2, c3, c4 = st.columns(4)
    ma = c1.text_input("Mã viên chức (*)", value=g("ma_vien_chuc"), disabled=edit)
    ho_ten = c2.text_input("Họ và tên (*)", value=g("ho_ten"))
    gioi_tinh = c3.selectbox("Giới tính", GIOI_TINH, index=pick(GIOI_TINH, g("gioi_tinh")))
    ngay_sinh = c4.date_input("Ngày sinh", value=d("ngay_sinh"), min_value=MIN_DATE, max_value=MAX_DATE, format=DATE_FORMAT)

    c5, c6, c7, c8 = st.columns(4)
    dan_toc = c5.text_input("Dân tộc", value=g("dan_toc"))
    cv_ids = list(CV_NAMES)
    cur_cv = info["chuc_vu_id"] if edit else None
    cv_idx = cv_ids.index(int(cur_cv)) if edit and pd.notna(cur_cv) and int(cur_cv) in cv_ids else 0
    chuc_vu_id = c6.selectbox("Chức vụ", options=cv_ids, index=cv_idx, format_func=CV_NAMES.get)
    ngay_td = c7.date_input("Ngày tuyển dụng", value=d("ngay_tuyen_dung"), min_value=MIN_DATE, max_value=MAX_DATE, format=DATE_FORMAT)
    ngay_dang = c8.date_input("Ngày vào Đảng", value=d("ngay_vao_dang"), min_value=MIN_DATE, max_value=MAX_DATE, format=DATE_FORMAT)

    st.markdown("**2. Thông tin Trình độ & Đào tạo**")
    c9, c10, c11, c12 = st.columns(4)
    trinh_do = c9.selectbox("Trình độ", TRINH_DO, index=pick(TRINH_DO, g("trinh_do")))
    truong_dt = c10.text_input("Tên trường đào tạo", value=g("ten_truong_dao_tao"))
    chuyen_nganh = c11.text_input("Chuyên ngành", value=g("chuyen_nganh"))
    ht_dt = c12.text_input("Hình thức đào tạo", value=g("hinh_thuc_dao_tao"))

    st.markdown("**3. Hạng chức danh & Các chứng chỉ**")
    c13, c14, c15 = st.columns(3)
    hang = c13.selectbox("Hạng CDNN", HANG_CDNN, index=pick(HANG_CDNN, g("hang_chuc_danh")))
    thac_si = c14.text_input("Thạc sĩ chuyên ngành", value=g("thac_si_chuyen_nganh"))
    qlnn = c15.text_input("Quản lý nhà nước", value=g("quan_ly_nha_nuoc"))

    c16, c17, c18, c19 = st.columns(4)
    llct = c16.text_input("Lý luận chính trị", value=g("ly_luan_chinh_tri"))
    tin_hoc = c17.text_input("Tin học", value=g("tin_hoc"))
    ngoai_ngu = c18.text_input("Ngoại ngữ", value=g("ngoai_ngu"))
    cc_khac = c19.text_input("Chứng chỉ khác", value=g("chung_chi_khac"))
    ghi_chu = st.text_input("Ghi chú thêm", value=g("ghi_chu"))

    vals = {
        "ho_ten": ho_ten.strip(), "gioi_tinh": gioi_tinh, "ngay_sinh": sql_date(ngay_sinh), "dan_toc": dan_toc,
        "chuc_vu_id": int(chuc_vu_id), "ngay_tuyen_dung": sql_date(ngay_td), "ngay_vao_dang": sql_date(ngay_dang),
        "hang_chuc_danh": hang, "trinh_do": trinh_do, "ten_truong_dao_tao": truong_dt, "chuyen_nganh": chuyen_nganh,
        "hinh_thuc_dao_tao": ht_dt, "thac_si_chuyen_nganh": thac_si, "quan_ly_nha_nuoc": qlnn,
        "ly_luan_chinh_tri": llct, "tin_hoc": tin_hoc, "ngoai_ngu": ngoai_ngu, "chung_chi_khac": cc_khac, "ghi_chu": ghi_chu,
    }
    return ma.strip(), vals

if menu == "📊 Tổng quan":
    st.title("📊 Bảng điều khiển Tổng quan")
    row = q("""SELECT COUNT(*) AS tong,
                      COALESCE(SUM(CASE WHEN gioi_tinh='Nam' THEN 1 ELSE 0 END), 0) AS nam,
                      COALESCE(SUM(CASE WHEN gioi_tinh='Nữ' THEN 1 ELSE 0 END), 0) AS nu FROM vien_chuc""").iloc[0]
    c1, c2, c3, c4 = st.columns(4)
    c1.metric("👥 Tổng số CB/GV", f"{int(row['tong'])} người")
    c2.metric("👨 Nam", f"{int(row['nam'])} người")
    c3.metric("👩 Nữ", f"{int(row['nu'])} người")
    c4.metric("🏆 Đang công tác", f"{int(row['tong'])} người")

elif menu == "👥 Quản lý Nhân sự":
    st.title("👥 Quản lý Hồ sơ Cán bộ, Giáo viên")
    if is_admin:
        tabs = st.tabs(["📋 Danh sách Đơn vị", "➕ Thêm mới", "✏️ Cập nhật & Xóa", "📂 Nhập Excel"])
        tab_list, tab_add, tab_edit, tab_import = tabs
    else:
        tab_list, tab_edit = st.tabs(["📋 Danh sách Đơn vị", "✏️ Cập nhật thông tin cá nhân"])

    with tab_list:
        st.subheader("Bảng thông tin Cán bộ - Giáo viên")
        df_vc = q("""
            SELECT vc.ma_vien_chuc AS "Mã VC", vc.ho_ten AS "Họ và tên", cv.ten_chuc_vu AS "Chức vụ",
                   vc.ngay_sinh AS "Ngày sinh", vc.gioi_tinh AS "Giới tính", vc.dan_toc AS "Dân tộc",
                   vc.ngay_tuyen_dung AS "Ngày tuyển dụng", vc.ngay_vao_dang AS "Ngày vào Đảng",
                   vc.hang_chuc_danh AS "Hạng CDNN", vc.trinh_do AS "Trình độ", vc.ten_truong_dao_tao AS "Trường đào tạo",
                   vc.chuyen_nganh AS "Chuyên ngành", vc.hinh_thuc_dao_tao AS "Hình thức ĐT", vc.thac_si_chuyen_nganh AS "Thạc sĩ",
                   vc.quan_ly_nha_nuoc AS "QL Nhà nước", vc.ly_luan_chinh_tri AS "LL Chính trị", vc.tin_hoc AS "Tin học",
                   vc.ngoai_ngu AS "Ngoại ngữ", vc.chung_chi_khac AS "CC Khác", vc.ghi_chu AS "Ghi chú"
            FROM vien_chuc vc LEFT JOIN chuc_vu cv ON vc.chuc_vu_id = cv.id
            ORDER BY vc.ma_vien_chuc ASC""")
        if not df_vc.empty:
            df_vc.insert(0, "STT", range(1, len(df_vc) + 1))
            for col in ["Ngày sinh", "Ngày tuyển dụng", "Ngày vào Đảng"]:
                df_vc[col] = df_vc[col].apply(format_date_vn)
            if is_admin:
                st.download_button("📥 Xuất danh sách nhân sự (Excel)", data=to_excel_bytes(df_vc, "NhanSu"),
                                   file_name=f"DanhSach_NhanSu_{date.today():%d%m%Y}.xlsx", mime=XLSX_MIME, type="primary")
            st.dataframe(df_vc, hide_index=True, **W)
        else:
            st.info("Chưa có dữ liệu.")

    if is_admin:
        with tab_add:
            with st.form("form_them_vc", clear_on_submit=True):
                ma_vc, vals = vc_fields()
                submitted = st.form_submit_button("💾 Thêm mới hồ sơ")
            if submitted:
                if ma_vc and vals["ho_ten"]:
                    try:
                        cols = ", ".join(["ma_vien_chuc"] + list(vals))
                        ph = ", ".join("?" * (len(vals) + 1))
                        run(f"INSERT INTO vien_chuc ({cols}) VALUES ({ph})", (ma_vc, *vals.values()))
                        flash_rerun("✅ Thêm mới thành công!", "nhan_su_editor")
                    except IntegrityError:
                        st.error("❌ Lỗi: Mã viên chức đã tồn tại!")
                else:
                    st.warning("⚠️ Vui lòng điền Mã VC và Họ tên (*)")

    with tab_edit:
        vc_edit_id = None
        if is_admin:
            st.write("Tick chọn 1 người để sửa hồ sơ, hoặc tick nhiều người để xóa:")
            df_ds = q("SELECT id, ma_vien_chuc, ho_ten FROM vien_chuc ORDER BY ma_vien_chuc ASC")
            if df_ds.empty:
                st.info("Chưa có hồ sơ.")
            else:
                df_ds.insert(0, "Chọn", False)
                ed = st.data_editor(df_ds, column_config={"Chọn": st.column_config.CheckboxColumn("Chọn", default=False), "id": None},
                                    disabled=["ma_vien_chuc", "ho_ten"], hide_index=True, key="nhan_su_editor", **W)
                sel = ed[ed["Chọn"]]
                if len(sel) == 1:
                    vc_edit_id = int(sel.iloc[0]["id"])
                elif len(sel) > 1:
                    ok = st.checkbox(f"Tôi xác nhận xóa {len(sel)} hồ sơ đã chọn (kèm chấm công, khen thưởng)", key="confirm_del_vc")
                    if st.button("❌ Xóa hàng loạt các hồ sơ đã chọn", disabled=not ok):
                        delete_vien_chuc(sel["id"].tolist())
                        flash_rerun("✅ Đã xóa hàng loạt thành công!", "nhan_su_editor", "confirm_del_vc")
        else:
            st.info("Bảng cập nhật thông tin cá nhân của bạn.")
            vc_edit_id = get_current_vc_id()
            if vc_edit_id is None: st.warning("Không tìm thấy hồ sơ gắn với tài khoản này. Vui lòng liên hệ Quản trị viên.")

        if vc_edit_id:
            found = q("SELECT * FROM vien_chuc WHERE id=?", (vc_edit_id,))
            if not found.empty:
                info = found.iloc[0]
                with st.form(f"form_sua_vc_{vc_edit_id}"):
                    _, vals = vc_fields(info)
                    col1, col2 = st.columns([1, 5])
                    btn_update = col1.form_submit_button("🔄 Lưu Cập nhật")
                    btn_delete = col2.form_submit_button("❌ Xóa Hồ sơ") if is_admin else False
                if btn_update:
                    if not vals["ho_ten"]:
                        st.warning("⚠️ Họ và tên không được để trống!")
                    else:
                        sets = ", ".join(f"{k}=?" for k in vals)
                        run(f"UPDATE vien_chuc SET {sets} WHERE id=?", (*vals.values(), vc_edit_id))
                        flash_rerun("✅ Cập nhật thành công!", "nhan_su_editor")
                if btn_delete:
                    delete_vien_chuc([vc_edit_id])
                    flash_rerun("✅ Đã xóa hồ sơ!", "nhan_su_editor")

    if is_admin:
        with tab_import:
            st.subheader("Nhập danh sách từ Excel")
            st.download_button("📥 Tải File Mẫu Excel", data=generate_excel_template(), file_name="Mau_DanhSach.xlsx", mime=XLSX_MIME)
            uploaded_file = st.file_uploader("Chọn file Excel đã điền", type=["xlsx"])
            if uploaded_file and st.button("🚀 Bắt đầu Import"):
                try:
                    df_excel = pd.read_excel(uploaded_file, dtype=str)
                    df_excel.columns = [str(c).strip() for c in df_excel.columns]
                    if "Mã VC" not in df_excel.columns or "Họ và tên" not in df_excel.columns:
                        st.error("File thiếu cột bắt buộc 'Mã VC' và 'Họ và tên'. Hãy dùng file mẫu.")
                    else:
                        cell = lambda r, k: safe_str(r.get(k, "")).strip()
                        records = []
                        for _, r in df_excel.iterrows():
                            ma, ten = cell(r, "Mã VC"), cell(r, "Họ và tên")
                            if not ma or not ten: continue
                            gt = "Nữ" if cell(r, "Giới tính").lower() in ("nữ", "nu", "f", "female") else "Nam"
                            records.append((ma, ten, gt, sql_date(date_or_none(r.get("Ngày sinh"))), cell(r, "Dân tộc"),
                                             sql_date(date_or_none(r.get("Ngày tuyển dụng"))),
                                             sql_date(date_or_none(r.get("Ngày vào Đảng"))),
                                             cell(r, "Trình độ"), cell(r, "Chuyên ngành"), cell(r, "Ghi chú")))
                        run_many("""INSERT INTO vien_chuc (ma_vien_chuc, ho_ten, gioi_tinh, ngay_sinh, dan_toc,
                                       ngay_tuyen_dung, ngay_vao_dang, trinh_do, chuyen_nganh, ghi_chu)
                                       VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?) ON CONFLICT (ma_vien_chuc) DO NOTHING""", records)
                        st.success(f"✅ Đã import / duyệt {len(records)} hồ sơ từ file!")
                except Exception as e:
                    st.error(f"Lỗi đọc file: {e}")

elif menu == "⏱️ Chấm công":
    st.title("⏱️ Bảng Chấm Công")
    today = date.today()
    col_m, col_y = st.columns(2)
    thang = col_m.selectbox("Chọn tháng", list(range(1, 13)), index=today.month - 1)
    nam = col_y.selectbox("Chọn năm", YEARS, index=pick(YEARS, today.year))

    last_day = calendar.monthrange(nam, thang)[1]
    days = [str(d) for d in range(1, last_day + 1)]
    start_date, end_date = f"{nam}-{thang:02d}-01", f"{nam}-{thang:02d}-{last_day:02d}"

    df_ds = q("SELECT id, ma_vien_chuc, ho_ten FROM vien_chuc ORDER BY ma_vien_chuc ASC")
    df_cc = q("""SELECT vien_chuc_id, CAST(strftime('%d', ngay_cham_cong) AS INTEGER) AS day, trang_thai
                 FROM cham_cong WHERE ngay_cham_cong BETWEEN ? AND ?""", (start_date, end_date))

    grid = pd.DataFrame("X", index=df_ds["id"].astype(int).tolist(), columns=days, dtype=object)
    for r in df_cc.itertuples(index=False):
        if int(r.vien_chuc_id) in grid.index and str(int(r.day)) in grid.columns:
            grid.at[int(r.vien_chuc_id), str(int(r.day))] = r.trang_thai
    df_display = df_ds.set_index("id").join(grid).reset_index()

    if df_ds.empty:
        st.warning("Chưa có dữ liệu cán bộ!")
    elif is_admin:
        tab_cc1, tab_cc2 = st.tabs(["📝 Chấm công Đơn vị", "📊 Thống kê ngày nghỉ"])
        with tab_cc1:
            col_config = {"id": None, "ma_vien_chuc": st.column_config.TextColumn("Mã VC", disabled=True), "ho_ten": st.column_config.TextColumn("Họ và tên", disabled=True)}
            for d in days: col_config[d] = st.column_config.SelectboxColumn(d, options=TRANG_THAI_CC, required=True, width="small")
            edited = st.data_editor(style_days(df_display, days), column_config=col_config,
                                    disabled=["id", "ma_vien_chuc", "ho_ten"], hide_index=True, height=500, key=f"cc_editor_{nam}_{thang}")
            if st.button("💾 LƯU BẢNG CHẤM CÔNG", type="primary"):
                records = [(int(row["id"]), f"{nam}-{thang:02d}-{int(d):02d}", row[d] or "X") for _, row in edited.iterrows() for d in days]
                run_many("""INSERT INTO cham_cong (vien_chuc_id, ngay_cham_cong, trang_thai) VALUES (?, ?, ?)
                            ON CONFLICT(vien_chuc_id, ngay_cham_cong) DO UPDATE SET trang_thai=EXCLUDED.trang_thai""", records)
                st.success("✅ Đã lưu chấm công!")
        with tab_cc2:
            df_tk = q("""SELECT vc.ma_vien_chuc AS "Mã VC", vc.ho_ten AS "Họ và tên", cc.trang_thai
                         FROM cham_cong cc JOIN vien_chuc vc ON cc.vien_chuc_id = vc.id
                         WHERE cc.ngay_cham_cong BETWEEN ? AND ? AND cc.trang_thai != 'X'""", (start_date, end_date))
            if not df_tk.empty:
                stat_df = df_tk.groupby(["Mã VC", "Họ và tên", "trang_thai"]).size().unstack(fill_value=0).reset_index()
                for tt in ["P", "K", "VR", "CT", "O"]:
                    if tt not in stat_df.columns: stat_df[tt] = 0
                st.dataframe(stat_df[["Mã VC", "Họ và tên", "P", "K", "VR", "CT", "O"]], hide_index=True, **W)
            else:
                st.success("Tháng này không có ngày nghỉ/công tác.")
    else:
        st.info("Chế độ xem bảng chấm công. Bạn không có quyền chỉnh sửa dữ liệu này.")
        view = df_display.drop(columns=["id"]).rename(columns={"ma_vien_chuc": "Mã VC", "ho_ten": "Họ và tên"})
        st.dataframe(style_days(view, days), hide_index=True, **W)

elif menu == "💰 Quản lý Lương":
    st.title("💰 Quản lý Lương Cán bộ, Giáo viên")

    def luong_form(vc_id, key):
        info = q("SELECT he_so_luong, phu_cap, ngay_huong_luong, chu_ky_nang_luong FROM vien_chuc WHERE id=?", (vc_id,)).iloc[0]
        with st.form(key):
            c1, c2, c3, c4 = st.columns(4)
            hs = c1.number_input("Hệ số lương", value=float(info["he_so_luong"]) if pd.notna(info["he_so_luong"]) else 0.0, min_value=0.0, step=0.33)
            pc = c2.number_input("Phụ cấp (hệ số)", value=float(info["phu_cap"]) if pd.notna(info["phu_cap"]) else 0.0, min_value=0.0, step=0.1)
            nhl = c3.date_input("Ngày hưởng hệ số", value=date_or_none(info["ngay_huong_luong"]), min_value=MIN_DATE, max_value=MAX_DATE, format=DATE_FORMAT)
            ck = c4.number_input("Chu kỳ nâng (Năm)", value=int(info["chu_ky_nang_luong"]) if pd.notna(info["chu_ky_nang_luong"]) else 3, min_value=1, step=1)
            submitted = st.form_submit_button("🔄 Cập nhật Lương")
        if submitted:
            run("UPDATE vien_chuc SET he_so_luong=?, phu_cap=?, ngay_huong_luong=?, chu_ky_nang_luong=? WHERE id=?", (hs, pc, sql_date(nhl), int(ck), int(vc_id)))
            flash_rerun("✅ Đã cập nhật lương!", "luong_editor_admin")

    if is_admin: tab_l1, tab_l2, tab_l_export = st.tabs(["📋 Bảng Lương Đơn vị", "✏️ Cập nhật Lương Đơn vị", "📂 Xuất DS Nâng Lương Quý"])
    else: tab_l1, tab_l2 = st.tabs(["📋 Bảng Lương Đơn vị", "✏️ Cập nhật Lương Cá nhân"])

    with tab_l1:
        df_luong = q("""
            SELECT vc.ma_vien_chuc AS "Mã VC", vc.ho_ten AS "Họ và tên", cv.ten_chuc_vu AS "Chức vụ",
                   COALESCE(vc.he_so_luong, 0) AS "Hệ số lương", COALESCE(vc.phu_cap, 0) AS "Phụ cấp",
                   vc.ngay_huong_luong AS "Ngày hưởng HS", COALESCE(vc.chu_ky_nang_luong, 3) AS "Chu kỳ (Năm)"
            FROM vien_chuc vc LEFT JOIN chuc_vu cv ON vc.chuc_vu_id = cv.id ORDER BY vc.ma_vien_chuc ASC""")
        if not df_luong.empty:
            df_luong.insert(0, "STT", range(1, len(df_luong) + 1))
            def calc_salary_info(row):
                tong = (row["Hệ số lương"] + row["Phụ cấp"]) * LUONG_CO_SO
                next_date, canh_bao, next_ts = "Chưa xác định", "⚪ Chưa có dữ liệu", pd.NaT
                nhl = date_or_none(row["Ngày hưởng HS"])
                if nhl and row["Chu kỳ (Năm)"] > 0:
                    next_ts = pd.Timestamp(nhl) + pd.DateOffset(years=int(row["Chu kỳ (Năm)"]))
                    next_date = next_ts.strftime("%d/%m/%Y")
                    left = (next_ts.date() - date.today()).days
                    canh_bao = "🔴 Đã quá hạn" if left < 0 else "🟡 Sắp đến hạn" if left <= 30 else "🟢 Bình thường"
                return {"Tổng lương": f"{tong:,.0f} đ", "Ngày nâng lương tới": next_date, "Cảnh báo": canh_bao, "Next_Obj": next_ts}
            extra = pd.DataFrame([calc_salary_info(r) for r in df_luong.to_dict("records")], index=df_luong.index)
            df_luong = pd.concat([df_luong, extra], axis=1)
            df_luong["Ngày hưởng HS"] = df_luong["Ngày hưởng HS"].apply(format_date_vn)
            df_luong_all = df_luong.copy()
            df_luong = df_luong.drop(columns=["Next_Obj"])
            st.dataframe(df_luong, hide_index=True, **W)
            if is_admin: st.download_button("📥 Xuất danh sách Lương", data=to_excel_bytes(df_luong, "Luong"), file_name="DS_Luong.xlsx", mime=XLSX_MIME)
        else:
            df_luong_all = pd.DataFrame()
            st.info("Chưa có dữ liệu.")

    with tab_l2:
        if is_admin:
            st.write("Tick chọn 1 cán bộ để cập nhật lương:")
            df_ds = q("SELECT id, ma_vien_chuc, ho_ten FROM vien_chuc ORDER BY ma_vien_chuc ASC")
            if df_ds.empty: st.info("Chưa có hồ sơ.")
            else:
                df_ds.insert(0, "Chọn", False)
                ed = st.data_editor(df_ds, column_config={"Chọn": st.column_config.CheckboxColumn("Chọn", default=False), "id": None},
                                    disabled=["ma_vien_chuc", "ho_ten"], hide_index=True, key="luong_editor_admin")
                sel = ed[ed["Chọn"]]
                if len(sel) == 1: luong_form(int(sel.iloc[0]["id"]), "form_sua_luong")
                elif len(sel) > 1: st.warning("Vui lòng chỉ chọn 1 cán bộ.")
        else:
            vc_id = get_current_vc_id()
            if vc_id:
                st.info("Cập nhật thông tin lương của cá nhân bạn.")
                luong_form(vc_id, "form_sua_luong_user")
            else: st.warning("Không tìm thấy hồ sơ gắn với tài khoản này.")

    if is_admin:
        with tab_l_export:
            quarter_export_ui(df_luong_all, "Next_Obj", ["Mã VC", "Họ và tên", "Chức vụ", "Hệ số lương", "Phụ cấp", "Ngày hưởng HS", "Chu kỳ (Năm)", "Ngày nâng lương tới"], "Báo cáo: Danh sách Cán bộ đến hạn Nâng lương", "DanhSach_NangLuong", "nl")

elif menu == "⏳ Quản lý Thâm niên":
    st.title("⏳ Quản lý Thâm niên Cán bộ, Giáo viên")
    st.info("Chỉ áp dụng cho Ban giám hiệu và Giáo viên.")

    if is_admin: tab_t1, tab_t_add, tab_t_edit, tab_t_export = st.tabs(["📋 Bảng Thâm niên Đơn vị", "➕ Thêm CBGV hưởng", "✏️ Cập nhật thông tin", "📂 Xuất DS Nâng Quý"])
    else: tab_t1, tab_t_edit = st.tabs(["📋 Bảng Thâm niên Đơn vị", "✏️ Cập nhật thông tin cá nhân"])

    df_tn_all = q("""
        SELECT vc.id, vc.ma_vien_chuc AS "Mã VC", vc.ho_ten AS "Họ và tên", cv.ten_chuc_vu AS "Chức vụ",
               vc.muc_tham_nien, vc.ngay_huong_tham_nien, vc.chu_ky_nang_tham_nien
        FROM vien_chuc vc LEFT JOIN chuc_vu cv ON vc.chuc_vu_id = cv.id
        WHERE COALESCE(cv.ten_chuc_vu, '') NOT LIKE ? ORDER BY vc.ma_vien_chuc ASC""", ('Nhân viên%',))

    def calc_tn(row):
        muc_v = row["muc_tham_nien"]
        muc = f"{muc_v:g}%" if pd.notna(muc_v) and muc_v > 0 else "Chưa hưởng"
        ngay = date_or_none(row["ngay_huong_tham_nien"])
        ck = int(row["chu_ky_nang_tham_nien"]) if pd.notna(row["chu_ky_nang_tham_nien"]) else 1
        next_s, canh_bao, next_ts = "Chưa xác định", "⚪ Chưa có dữ liệu", pd.NaT
        if ngay:
            next_ts = pd.Timestamp(ngay) + pd.DateOffset(years=ck)
            next_s = next_ts.strftime("%d/%m/%Y")
            left = (next_ts.date() - date.today()).days
            canh_bao = "🔴 Đã quá hạn" if left < 0 else "🟡 Sắp đến hạn" if left <= 30 else "🟢 Bình thường"
        return {"Mức hiện hưởng": muc, "Ngày hưởng": format_date_vn(ngay), "Chu kỳ (Năm)": ck, "Ngày hưởng lần sau": next_s, "Cảnh báo": canh_bao, "Next_Obj": next_ts}

    if not df_tn_all.empty:
        extra = pd.DataFrame([calc_tn(r) for r in df_tn_all.to_dict("records")], index=df_tn_all.index)
        df_tn_all = pd.concat([df_tn_all, extra], axis=1)
        df_da_huong = df_tn_all[df_tn_all["muc_tham_nien"] > 0].copy()
    else: df_da_huong = pd.DataFrame()

    def tn_form(info, key, submit_label):
        with st.form(key):
            c1, c2, c3 = st.columns(3)
            muc = c1.number_input("Mức hiện hưởng (%)", value=float(info["muc_tham_nien"]) if pd.notna(info["muc_tham_nien"]) else 5.0, min_value=0.0, step=1.0)
            ng = c2.date_input("Ngày hưởng thâm niên", value=date_or_today(info["ngay_huong_tham_nien"]), min_value=MIN_DATE, max_value=MAX_DATE, format=DATE_FORMAT)
            ck = c3.number_input("Chu kỳ nâng (Năm)", value=int(info["chu_ky_nang_tham_nien"]) if pd.notna(info["chu_ky_nang_tham_nien"]) else 1, min_value=1, step=1)
            submitted = st.form_submit_button(submit_label)
        if submitted:
            run("UPDATE vien_chuc SET muc_tham_nien=?, ngay_huong_tham_nien=?, chu_ky_nang_tham_nien=? WHERE id=?", (muc, sql_date(ng), int(ck), int(info["id"])))
            flash_rerun("✅ Đã cập nhật!", "tn_editor_admin")

    with tab_t1:
        if not df_da_huong.empty:
            df_show = df_da_huong.drop(columns=["id", "muc_tham_nien", "ngay_huong_tham_nien", "chu_ky_nang_tham_nien", "Next_Obj"])
            df_show.insert(0, "STT", range(1, len(df_show) + 1))
            st.dataframe(df_show, hide_index=True, **W)
        else: st.info("Chưa có danh sách hưởng thâm niên.")

    if is_admin:
        with tab_t_add:
            df_chua = df_tn_all[df_tn_all["muc_tham_nien"].isna() | (df_tn_all["muc_tham_nien"] <= 0)] if not df_tn_all.empty else pd.DataFrame()
            if not df_chua.empty:
                chua_dict = {int(i): f"{m} - {t}" for i, m, t in zip(df_chua["id"], df_chua["Mã VC"], df_chua["Họ và tên"])}
                cb_id = st.selectbox("Chọn Cán bộ - Giáo viên", options=list(chua_dict), format_func=chua_dict.get)
                with st.form("form_them_moi_tn"):
                    c1, c2, c3 = st.columns(3)
                    muc_add = c1.number_input("Mức thâm niên ban đầu (%)", value=5.0, min_value=0.0, step=1.0)
                    ng_add = c2.date_input("Ngày bắt đầu hưởng", value=date.today(), min_value=MIN_DATE, max_value=MAX_DATE, format=DATE_FORMAT)
                    ck_add = c3.number_input("Chu kỳ nâng tiếp theo (Năm)", value=1, min_value=1, step=1)
                    submitted = st.form_submit_button("✅ Đưa vào danh sách hưởng")
                if submitted:
                    run("UPDATE vien_chuc SET muc_tham_nien=?, ngay_huong_tham_nien=?, chu_ky_nang_tham_nien=? WHERE id=?", (muc_add, sql_date(ng_add), int(ck_add), int(cb_id)))
                    flash_rerun("✅ Đã thêm thành công!", "tn_editor_admin")
            else: st.success("Tất cả CBGV thuộc diện hưởng đã được cấu hình Thâm niên!")

    with tab_t_edit:
        if is_admin:
            if not df_da_huong.empty:
                st.write("Tick chọn 1 người để sửa:")
                df_pick = df_da_huong[["id", "Mã VC", "Họ và tên"]].copy()
                df_pick.insert(0, "Chọn", False)
                ed = st.data_editor(df_pick, column_config={"Chọn": st.column_config.CheckboxColumn("Chọn", default=False), "id": None}, disabled=["Mã VC", "Họ và tên"], hide_index=True, key="tn_editor_admin")
                sel = ed[ed["Chọn"]]
                if len(sel) == 1: tn_form(df_da_huong[df_da_huong["id"] == int(sel.iloc[0]["id"])].iloc[0], "form_sua_tn", "🔄 Cập nhật")
                elif len(sel) > 1: st.warning("Vui lòng chỉ chọn 1 người.")
            else: st.info("Chưa có ai trong danh sách hưởng thâm niên.")
        else:
            vc_id = get_current_vc_id()
            mine = df_da_huong[df_da_huong["id"] == vc_id] if vc_id and not df_da_huong.empty else pd.DataFrame()
            if not mine.empty:
                st.info("Cập nhật thông tin thâm niên của cá nhân bạn.")
                tn_form(mine.iloc[0], "form_sua_tn_user", "🔄 Cập nhật Thâm niên")
            else: st.warning("Bạn chưa có dữ liệu hưởng thâm niên, hoặc chưa đủ điều kiện.")

    if is_admin:
        with tab_t_export:
            quarter_export_ui(df_da_huong, "Next_Obj", ["Mã VC", "Họ và tên", "Chức vụ", "Mức hiện hưởng", "Ngày hưởng", "Chu kỳ (Năm)", "Ngày hưởng lần sau"], "Báo cáo: Danh sách Cán bộ đến hạn Nâng thâm niên", "DanhSach_NangThamNien", "tn")

elif menu == "🏆 Thi đua - Khen thưởng":
    st.title("🏆 Quản lý Thi đua - Khen thưởng")
    tab_k1, tab_k2, tab_k3 = st.tabs(["📋 Danh sách", "➕ Thêm mới", "✏️ Cập nhật & Xóa"])
    where, params = scope_where("vc")
    df_kt = q(f"""
        SELECT kt.id, vc.ma_vien_chuc AS "Mã VC", vc.ho_ten AS "Họ và tên",
               kt.ngay_thang AS "Ngày tháng", kt.thanh_tich AS "Thành tích",
               kt.danh_hieu AS "Danh hiệu", kt.hinh_thuc AS "Hình thức",
               kt.cap_khen AS "Cấp khen", kt.nam_khen AS "Năm", kt.so_quyet_dinh AS "Số QĐ"
        FROM khen_thuong kt JOIN vien_chuc vc ON kt.vien_chuc_id = vc.id {where} ORDER BY kt.ngay_thang DESC""", params)

    with tab_k1:
        if not df_kt.empty:
            df_show = df_kt.drop(columns=["id"])
            df_show["Ngày tháng"] = df_show["Ngày tháng"].apply(format_date_vn)
            df_show["Năm"] = df_show["Năm"].astype("Int64")
            df_show.insert(0, "STT", range(1, len(df_show) + 1))
            st.dataframe(df_show, hide_index=True, **W)
        else: st.info("Chưa có hồ sơ khen thưởng.")

    with tab_k2:
        w2, p2 = scope_where("vien_chuc")
        df_vc_add = q(f"SELECT id, ma_vien_chuc, ho_ten FROM vien_chuc {w2} ORDER BY ma_vien_chuc", p2)
        if not df_vc_add.empty:
            names = {int(i): f"{m} - {t}" for i, m, t in zip(df_vc_add["id"], df_vc_add["ma_vien_chuc"], df_vc_add["ho_ten"])}
            with st.form("form_them_kt", clear_on_submit=True):
                c1, c2 = st.columns(2)
                if is_admin: vc_id = c1.selectbox("Chọn Cán bộ", options=list(names), format_func=names.get)
                else:
                    vc_id = int(df_vc_add.iloc[0]["id"])
                    c1.text_input("Tên Cán bộ", value=df_vc_add.iloc[0]["ho_ten"], disabled=True)
                ngay_thang = c2.date_input("Ngày tháng", value=date.today(), min_value=MIN_DATE, max_value=MAX_DATE, format=DATE_FORMAT)
                c3, c4 = st.columns(2)
                thanh_tich = c3.text_input("Thành tích")
                danh_hieu = c4.text_input("Danh hiệu")
                c5, c6 = st.columns(2)
                hinh_thuc = c5.text_input("Hình thức khen")
                cap_khen = c6.text_input("Cấp khen")
                c7, c8 = st.columns(2)
                nam_khen = c7.number_input("Năm được khen", min_value=1970, max_value=2100, value=date.today().year, step=1)
                so_qd = c8.text_input("Số quyết định")
                submitted = st.form_submit_button("💾 Thêm khen thưởng")
            if submitted:
                run("""INSERT INTO khen_thuong (vien_chuc_id, ngay_thang, thanh_tich, danh_hieu, hinh_thuc, cap_khen, nam_khen, so_quyet_dinh)
                       VALUES (?, ?, ?, ?, ?, ?, ?, ?)""", (int(vc_id), sql_date(ngay_thang), thanh_tich, danh_hieu, hinh_thuc, cap_khen, int(nam_khen), so_qd))
                flash_rerun("✅ Đã thêm khen thưởng thành công!")
        else: st.warning("Chưa có hồ sơ cán bộ nào hợp lệ!")

    with tab_k3:
        if not df_kt.empty:
            st.write("Chọn thành tích khen thưởng để sửa đổi:")
            kt_dict = {int(r["id"]): f"{'' if pd.isna(r['Năm']) else int(r['Năm'])} - {r['Họ và tên']} ({safe_str(r['Thành tích'])})" for r in df_kt.to_dict("records")}
            kt_id = st.selectbox("Danh mục Khen thưởng đã ghi nhận", options=list(kt_dict), format_func=kt_dict.get, index=None)
            if kt_id is not None:
                kt_info = q("SELECT * FROM khen_thuong WHERE id=?", (int(kt_id),)).iloc[0]
                with st.form(f"form_sua_kt_{kt_id}"):
                    st.write(f"Đang sửa: **{kt_dict[kt_id]}**")
                    ng = st.date_input("Ngày tháng", value=date_or_today(kt_info["ngay_thang"]), min_value=MIN_DATE, max_value=MAX_DATE, format=DATE_FORMAT)
                    c1, c2 = st.columns(2)
                    tt = c1.text_input("Thành tích", value=safe_str(kt_info["thanh_tich"]))
                    dh = c2.text_input("Danh hiệu", value=safe_str(kt_info["danh_hieu"]))
                    ht = c1.text_input("Hình thức khen", value=safe_str(kt_info["hinh_thuc"]))
                    ck = c2.text_input("Cấp khen", value=safe_str(kt_info["cap_khen"]))
                    nk = c1.number_input("Năm được khen", min_value=1970, max_value=2100, step=1, value=int(kt_info["nam_khen"]) if pd.notna(kt_info["nam_khen"]) else date.today().year)
                    qd = c2.text_input("Số quyết định", value=safe_str(kt_info["so_quyet_dinh"]))
                    b1, b2 = st.columns([1, 5])
                    do_update = b1.form_submit_button("🔄 Lưu Cập nhật")
                    do_delete = b2.form_submit_button("❌ Xóa Khen thưởng")
                if do_update:
                    run("""UPDATE khen_thuong SET ngay_thang=?, thanh_tich=?, danh_hieu=?, hinh_thuc=?, cap_khen=?, nam_khen=?, so_quyet_dinh=? WHERE id=?""", (sql_date(ng), tt, dh, ht, ck, int(nk), qd, int(kt_id)))
                    flash_rerun("✅ Cập nhật thành công!")
                if do_delete:
                    run("DELETE FROM khen_thuong WHERE id=?", (int(kt_id),))
                    flash_rerun("✅ Đã xóa khen thưởng!")
        else: st.info("Không có dữ liệu để sửa.")

elif menu == "🚨 Quản lý Vụ việc":
    st.title("🚨 Quản lý Vụ việc")
    if is_admin: tab_v1, tab_v2, tab_v3 = st.tabs(["📋 Danh sách Vụ việc", "➕ Thêm Vụ việc", "✏️ Cập nhật & Xóa"])
    else: (tab_v1,) = st.tabs(["📋 Danh sách Vụ việc"])

    df_vv = q("""SELECT id, ngay_thang AS "Ngày tháng", noi_dung AS "Nội dung vụ việc", phuong_an AS "Phương án xử lý" FROM vu_viec ORDER BY ngay_thang DESC""")
    with tab_v1:
        if not df_vv.empty:
            df_show = df_vv.drop(columns=["id"])
            df_show["Ngày tháng"] = df_show["Ngày tháng"].apply(format_date_vn)
            df_show.insert(0, "STT", range(1, len(df_show) + 1))
            st.dataframe(df_show, hide_index=True, **W)
        else: st.info("Chưa có vụ việc nào được ghi nhận.")

    if is_admin:
        with tab_v2:
            with st.form("form_them_vv", clear_on_submit=True):
                ngay_vv = st.date_input("Ngày tháng", value=date.today(), min_value=MIN_DATE, max_value=MAX_DATE, format=DATE_FORMAT)
                noi_dung = st.text_area("Nội dung vụ việc (*)")
                phuong_an = st.text_area("Phương án xử lý")
                submitted = st.form_submit_button("💾 Lưu Vụ việc")
            if submitted:
                if noi_dung.strip():
                    run("INSERT INTO vu_viec (ngay_thang, noi_dung, phuong_an) VALUES (?, ?, ?)", (sql_date(ngay_vv), noi_dung, phuong_an))
                    flash_rerun("✅ Đã thêm vụ việc thành công!", "vv_editor")
                else: st.warning("Vui lòng điền nội dung vụ việc!")
        with tab_v3:
            if not df_vv.empty:
                st.write("Tick chọn để sửa:")
                df_pick = df_vv[["id", "Ngày tháng", "Nội dung vụ việc"]].copy()
                df_pick["Ngày tháng"] = df_pick["Ngày tháng"].apply(format_date_vn)
                df_pick.insert(0, "Chọn", False)
                ed = st.data_editor(df_pick, column_config={"Chọn": st.column_config.CheckboxColumn("Chọn", default=False), "id": None}, disabled=["Ngày tháng", "Nội dung vụ việc"], hide_index=True, key="vv_editor")
                sel = ed[ed["Chọn"]]
                if len(sel) == 1:
                    vv_id = int(sel.iloc[0]["id"])
                    vv = q("SELECT * FROM vu_viec WHERE id=?", (vv_id,)).iloc[0]
                    with st.form(f"form_sua_vv_{vv_id}"):
                        ng = st.date_input("Ngày tháng", value=date_or_today(vv["ngay_thang"]), min_value=MIN_DATE, max_value=MAX_DATE, format=DATE_FORMAT)
                        nd = st.text_area("Nội dung vụ việc", value=safe_str(vv["noi_dung"]))
                        pa = st.text_area("Phương án xử lý", value=safe_str(vv["phuong_an"]))
                        b1, b2 = st.columns([1, 5])
                        do_update = b1.form_submit_button("🔄 Lưu Cập nhật")
                        do_delete = b2.form_submit_button("❌ Xóa Vụ việc")
                    if do_update:
                        run("UPDATE vu_viec SET ngay_thang=?, noi_dung=?, phuong_an=? WHERE id=?", (sql_date(ng), nd, pa, vv_id))
                        flash_rerun("✅ Cập nhật thành công!", "vv_editor")
                    if do_delete:
                        run("DELETE FROM vu_viec WHERE id=?", (vv_id,))
                        flash_rerun("✅ Đã xóa vụ việc!", "vv_editor")
                elif len(sel) > 1: st.warning("Vui lòng chỉ chọn 1 vụ việc.")
            else: st.info("Không có dữ liệu để sửa.")

elif menu == "⚙️ Quản lý Tài khoản" and is_admin:
    st.title("⚙️ Quản lý Tài khoản")
    tab_a1, tab_a2 = st.tabs(["➕ Tạo tài khoản", "📋 Danh sách & Cập nhật/Xóa"])
    df_all_vc = q("SELECT ma_vien_chuc, ho_ten FROM vien_chuc ORDER BY ma_vien_chuc")
    vc_names = dict(zip(df_all_vc["ma_vien_chuc"], df_all_vc["ho_ten"]))
    ma_options = [""] + list(vc_names)
    ma_label = lambda x: "Không liên kết" if x == "" else f"{x} - {vc_names.get(x, '')}"
    ROLE_LABEL = {"user": "Giáo viên / Cán bộ", "admin": "Quản trị viên (Admin)"}
    LABEL_ROLE = {v: k for k, v in ROLE_LABEL.items()}

    with tab_a1:
        with st.form("form_create_user", clear_on_submit=True):
            st.subheader("Tạo tài khoản mới")
            c1, c2 = st.columns(2)
            new_user = c1.text_input("Tên đăng nhập (*)")
            new_pass = c2.text_input("Mật khẩu (*)", type="password")
            new_role = c1.selectbox("Quyền", options=list(LABEL_ROLE))
            ma_link = c2.selectbox("Liên kết với Cán bộ (Bắt buộc với Giáo viên)", options=ma_options, format_func=ma_label)
            submitted = st.form_submit_button("💾 Tạo tài khoản")
        if submitted:
            role = LABEL_ROLE[new_role]
            if not (new_user.strip() and new_pass): st.warning("⚠️ Vui lòng điền đủ Tên đăng nhập và Mật khẩu!")
            elif role == "user" and not ma_link: st.warning("⚠️ Tài khoản Giáo viên phải được liên kết với một hồ sơ Cán bộ!")
            else:
                try:
                    run("INSERT INTO users (username, password, role, ma_vien_chuc) VALUES (?, ?, ?, ?)", (new_user.strip(), make_hashes(new_pass), role, ma_link))
                    flash_rerun(f"✅ Đã tạo tài khoản: {new_user.strip()}", "user_editor")
                except IntegrityError: st.error("❌ Tên đăng nhập đã tồn tại!")

    with tab_a2:
        df_users = q("SELECT id, username, role, ma_vien_chuc FROM users ORDER BY username")
        show = pd.DataFrame({"Chọn": False, "id": df_users["id"], "username": df_users["username"], "Quyền": df_users["role"].map({"user": "Giáo viên / Cán bộ", "admin": "Quản trị viên"}), "ma_vien_chuc": df_users["ma_vien_chuc"].fillna("")})
        st.write("Tick chọn 1 tài khoản để cấu hình hoặc xóa:")
        ed = st.data_editor(show, column_config={"Chọn": st.column_config.CheckboxColumn("Chọn", default=False), "id": None, "username": "Tài khoản", "ma_vien_chuc": "Mã VC Liên kết"}, disabled=["username", "Quyền", "ma_vien_chuc"], hide_index=True, key="user_editor", **W)
        sel = ed[ed["Chọn"]]
        if len(sel) == 1:
            u_id = int(sel.iloc[0]["id"])
            u = df_users[df_users["id"] == u_id].iloc[0]
            is_root = u["username"] == "admin"
            is_self = u["username"] == st.session_state["username"]
            with st.form(f"form_edit_user_{u_id}"):
                st.markdown(f"### Cập nhật: `{u['username']}`")
                c1, c2 = st.columns(2)
                pw_new = c1.text_input("Mật khẩu mới (Để trống nếu không muốn đổi)", type="password")
                role_names = list(LABEL_ROLE)
                role_new = c2.selectbox("Quyền", role_names, index=pick(role_names, ROLE_LABEL.get(u["role"], "")), disabled=is_root)
                ma_new = st.selectbox("Liên kết với Cán bộ", options=ma_options, index=pick(ma_options, safe_str(u["ma_vien_chuc"])), format_func=ma_label)
                b1, b2 = st.columns([1, 5])
                do_update = b1.form_submit_button("🔄 Lưu Cập nhật")
                do_delete = b2.form_submit_button("❌ Xóa Tài khoản")
            if do_update:
                role_val = "admin" if is_root else LABEL_ROLE[role_new]
                if role_val == "user" and not ma_new: st.warning("⚠️ Tài khoản Giáo viên phải được liên kết với một hồ sơ Cán bộ!")
                elif is_self and role_val != "admin": st.error("⚠️ Bạn không thể tự hạ quyền của chính mình!")
                else:
                    if pw_new.strip(): run("UPDATE users SET password=?, role=?, ma_vien_chuc=? WHERE id=?", (make_hashes(pw_new), role_val, ma_new, u_id))
                    else: run("UPDATE users SET role=?, ma_vien_chuc=? WHERE id=?", (role_val, ma_new, u_id))
                    flash_rerun("✅ Cập nhật tài khoản thành công!", "user_editor")
            if do_delete:
                if is_root: st.error("⚠️ Không thể xóa tài khoản Quản trị viên gốc (admin)!")
                elif is_self: st.error("⚠️ Bạn không thể tự xóa tài khoản đang đăng nhập!")
                else:
                    run("DELETE FROM users WHERE id=?", (u_id,))
                    flash_rerun("✅ Đã xóa tài khoản!", "user_editor")
        elif len(sel) > 1: st.warning("Vui lòng chỉ chọn 1 tài khoản.")
